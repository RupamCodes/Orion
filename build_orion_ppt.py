#!/usr/bin/env python3
import os
import shutil
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_PATH = Path("/Users/rupamdolui/Orion/Orion/ORION_Updated_Review.pptx")

# Exact replacements mapped from ORION_Review_2_Final.pptx
REPLACEMENTS = {
    # Slide 4: Abstract
    "The current implementation combines a Next.js/React interface with a Python FastAPI service and optional local Ollama inference. Documents are normalized, chunked, indexed in memory, retrieved with source filtering and term-frequency scoring, and passed to a grounded response layer.":
        "The implementation combines a Next.js 14 / React 18 interface with a Python FastAPI backend, SQLite persistence (SQLAlchemy 2.0), ChromaDB vector storage, and local Ollama (llama3.2) inference. Documents are validated, chunked with section/position metadata, and retrieved via hybrid semantic (all-MiniLM-L6-v2) and BM25 scoring.",
    "Top-K retrieval with citations":
        "Hybrid Semantic + BM25 retrieval with citations",
    "Study and Studio artifacts":
        "5-Tool Studio Forge & SQLite persistence",

    # Slide 5: Introduction
    "A single workspace for document ingestion, grounded Q&A, citations, notes and study-oriented synthesis.":
        "A persistent multi-notebook workspace for ingestion, hybrid retrieval, grounded Q&A, citations, notes, and 5 Studio synthesis tools.",
    "The application can use Ollama at localhost:11434 and also contains a browser-side retrieval path.":
        "Runs Ollama (llama3.2) at localhost:11434 with local all-MiniLM-L6-v2 embeddings in ChromaDB, SQLite persistence, and deterministic fallbacks.",

    # Slide 6: Problem Statement
    "Current work covers ingestion, chunking, retrieval, grounded Q&A, source viewing, notes and Studio synthesis.":
        "Covers SQLite multi-notebook persistence, hybrid ChromaDB + BM25 retrieval, grounded Ollama chat with citations, and 5 Studio tools.",
    "No verified production-scale benchmark or deployment performance claim is included in the project files.":
        "Validated via 99 pytest backend tests, 6 batch verification checks, and 10/10 live API audits; cloud hosting is out of scope.",

    # Slide 7: Literature Survey
    "ORION currently uses lightweight term-frequency retrieval rather than DPR.":
        "ORION implements dense embeddings (all-MiniLM-L6-v2 in ChromaDB) combined with BM25 in a hybrid retrieval pipeline.",
    "Survey-level taxonomy; ORION implements a practical local pipeline.":
        "Survey-level taxonomy; ORION implements a persistent local hybrid RAG and Studio pipeline.",

    # Slide 10: Proposed System
    "Accept document text and supported source formats; normalize content before indexing.":
        "Validate and persist multi-source documents per isolated notebook in SQLite and ChromaDB.",
    "Split content into bounded passages while retaining source and section metadata.":
        "Split content into bounded passages retaining source, section, and document-position metadata.",
    "Filter by selected sources and rank candidate chunks using term-frequency scoring.":
        "Filter by selected sources and rank chunks using hybrid semantic (all-MiniLM-L6-v2) + BM25 scoring.",
    "Use grounded context with local Ollama when available; otherwise use local extraction fallback.":
        "Generate grounded answers via local Ollama (llama3.2) with deterministic source-grounded fallback.",
    "Expose source, chunk, snippet and page metadata alongside the answer.":
        "Expose clickable citation chips, source IDs, snippets, and integrated source-viewer navigation.",
    "Create notes, briefings, mind maps, flashcards and podcast-oriented artifacts.":
        "Generate persisted Audio Overviews (TTS), Flashcards, Mind Maps, Reports, and SVG Infographics.",

    # Slide 11: Architecture Diagram
    "450-char chunks\n80-char overlap":
        "Section & position\nmetadata chunks",
    "In-memory\nsource/chunk store":
        "SQLite + ChromaDB\n(all-MiniLM-L6-v2)",
    "Selected sources\nTop-K scoring":
        "Hybrid Semantic\n+ BM25 Top-K",
    "Ollama\nor local fallback":
        "Ollama (llama3.2)\n+ offline fallback",
    "Backend: FastAPI + Python":
        "Backend: FastAPI + SQLAlchemy 2.0 (SQLite) + ChromaDB",
    "Local inference: Ollama  ·  PDF parsing: pdf.js / pdf-parse  ·  OCR dependency: Tesseract.js":
        "Local AI: Ollama (llama3.2)  ·  Embeddings: all-MiniLM-L6-v2  ·  Retrieval: Hybrid Semantic + BM25",

    # Slide 12: Requirements
    "Node.js / npm\nNext.js 14 · React 18 · TypeScript 5.7\nPython · FastAPI · Uvicorn\nOllama (optional local inference)":
        "Next.js 14.2.23 · React 18 · TypeScript · Tailwind\nPython · FastAPI · SQLAlchemy 2.0 · SQLite\nChromaDB · Sentence-Transformers (all-MiniLM-L6-v2)\nOllama (llama3.2) & Browser Web Speech TTS",

    # Slide 13: Working Flow
    "Chunk + metadata": "Chunk + ChromaDB",
    "Retrieve Top-K": "Hybrid RAG Top-K",
    "Show citations / save notes": "Cite / Studio Forge",
    "If Ollama is unavailable, the frontend contains a direct answer-extraction fallback from retrieved passages.":
        "All notebooks, chats, notes, and Studio artifacts persist in SQLite; deterministic fallback protects offline generation.",
    "Notes can also be converted back into first-class sources for recursive retrieval in the Python API.":
        "Studio Forge generates persisted Audio Overviews (dual-host TTS), Flashcards, Mind Maps, Reports, and SVG Infographics.",

    # Slide 14: Results
    "Multi-source workspace with source selection and source viewing.":
        "Persistent SQLite multi-notebook workspace with source validation, viewing, and vector cleanup.",
    "Top-K retrieved passages are passed into a grounded answer layer.":
        "Hybrid semantic (all-MiniLM-L6-v2) + BM25 retrieval passed to Ollama (llama3.2) grounded chat.",
    "Citation objects retain source ID, title, chunk/page metadata and snippets.":
        "Interactive citation chips navigate directly to source passages with safe Markdown rendering.",
    "Backend endpoints support briefing, mind map, flashcards and podcast artifacts.":
        "5 persisted tools: Audio Overview (TTS), Flashcards, Mind Map, Reports, and Infographics (SVG).",
    "Frontend checks localhost Ollama and streams model responses when available.":
        "Live Ollama health status, dark/light theme, collapsible panels, and resilient audio fallback.",
    "Pinned notes can be converted into first-class indexed sources.":
        "Verified: 99/99 pytest tests, 6/6 batch4 tests, 10/10 live API audits, and clean Next.js build.",
    "No live UI screenshots were supplied; this slide deliberately avoids fabricated performance metrics or screenshots.":
        "All persistence, hybrid retrieval, chat, and Studio Forge endpoints verified via automated test suites.",

    # Slide 15: Future Enhancement
    "Stronger semantic retrieval":
        "Cross-encoder reranking & OCR",
    "Replace lightweight term-frequency ranking with embedding-based retrieval and a persistent vector store.":
        "Extend the ChromaDB + BM25 hybrid pipeline with cross-encoder reranking and complex table/OCR extraction.",
    "Add a repeatable QA benchmark with retrieval precision/recall, answer faithfulness and latency measurements.":
        "Add automated Playwright visual verification and quantitative RAG precision/faithfulness benchmarks.",

    # Slide 16: Conclusion
    "Built with Next.js/React and a Python FastAPI service.":
        "Built with Next.js 14, React 18, FastAPI, SQLite (SQLAlchemy 2.0), and ChromaDB.",
    "Supports source-aware chunking, Top-K retrieval and citation metadata.":
        "Implements hybrid semantic (all-MiniLM-L6-v2) + BM25 retrieval and clickable citations.",
    "Integrates optional local Ollama inference without requiring a cloud API key.":
        "Delivers local Ollama (llama3.2) chat and 5 persisted Studio Forge synthesis tools.",
    "Provides a foundation for stronger semantic retrieval and measurable evaluation.":
        "Validated across 99 pytest backend tests, 6 batch verification checks, and 10 live API audits.",
}

def replace_in_text_frame(tf):
    full_text = tf.text.strip()
    for old, new in REPLACEMENTS.items():
        if old in full_text:
            # Preserve first run's font styling if updating a single-paragraph frame
            if len(tf.paragraphs) == 1 and tf.paragraphs[0].runs:
                p = tf.paragraphs[0]
                r0 = p.runs[0]
                fname, fsize, fbold, fcolor = r0.font.name, r0.font.size, r0.font.bold, (r0.font.color.rgb if r0.font.color and r0.font.color.type == 1 else None)
                updated = full_text.replace(old, new)
                p.text = updated
                if p.runs:
                    p.runs[0].font.name = fname
                    p.runs[0].font.size = fsize
                    p.runs[0].font.bold = fbold
                    if fcolor:
                        p.runs[0].font.color.rgb = fcolor
                return
        # Also check paragraph-by-paragraph
        for p in tf.paragraphs:
            p_text = p.text.strip()
            if old in p_text and p.runs:
                r0 = p.runs[0]
                fname, fsize, fbold, fcolor = r0.font.name, r0.font.size, r0.font.bold, (r0.font.color.rgb if r0.font.color and r0.font.color.type == 1 else None)
                p.text = p.text.replace(old, new)
                if p.runs:
                    p.runs[0].font.name = fname
                    p.runs[0].font.size = fsize
                    p.runs[0].font.bold = fbold
                    if fcolor:
                        p.runs[0].font.color.rgb = fcolor

def find_template():
    candidates = [
        Path("/Users/rupamdolui/Orion/Orion/ORION_Review_2_Final.pptx"),
        Path("/Users/rupamdolui/Orion/ORION_Review_2_Final.pptx"),
        Path.home() / "Downloads" / "ORION_Review_2_Final.pptx",
        Path.home() / "Desktop" / "ORION_Review_2_Final.pptx",
    ]
    for c in candidates:
        if c.exists():
            return c
    return None

def update_existing_template(template_path: Path):
    prs = Presentation(str(template_path))
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                replace_in_text_frame(shape.text_frame)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        replace_in_text_frame(cell.text_frame)
            if shape.shape_type == MSO_SHAPE.GROUP:
                for sub in shape.shapes:
                    if sub.has_text_frame:
                        replace_in_text_frame(sub.text_frame)
    prs.save(str(OUTPUT_PATH))
    print(f"Updated directly from template ({template_path}) -> {OUTPUT_PATH}")

def build_from_scratch():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    BG = RGBColor(13, 17, 23)
    CARD = RGBColor(22, 27, 34)
    BORDER = RGBColor(48, 54, 61)
    CYAN = RGBColor(0, 245, 212)
    WHITE = RGBColor(240, 246, 252)
    MUTED = RGBColor(139, 148, 158)

    slides_data = [
        ("01", "ORION — Local AI Research Workspace & RAG", [
            "Review 2 / 3 · Document Intelligence & Grounded Research",
            "Presented by: Prasiddha Raj Sedai (24BBTCA089), Rupam Dolui (24BBTCA098), Satyam Mohanty (24BBTCA107)",
            "Ask. Retrieve. Understand."
        ]),
        ("02", "Team Members", [
            "001 · Prasiddha Raj Sedai (24BBTCA089) — Project Contributor",
            "002 · Rupam Dolui (24BBTCA098) — Project Contributor",
            "003 · Satyam Mohanty (24BBTCA107) — Project Contributor"
        ]),
        ("03", "Agenda", [
            "01 Abstract   |   02 Introduction   |   03 Problem Statement   |   04 Literature Survey",
            "05 Existing System   |   06 Issues in Existing System   |   07 Proposed System",
            "08 Architecture   |   09 Requirements   |   10 Working Flow   |   11 Results",
            "12 Future Enhancement   |   13 Conclusion   |   14 References"
        ]),
        ("04", "Abstract", [
            "ORION is a local AI research workspace that lets users load documents into persistent notebooks, retrieve relevant passages, ask grounded questions, and inspect source citations.",
            "Combines Next.js 14 / React 18 with Python FastAPI, SQLite (SQLAlchemy 2.0), ChromaDB (all-MiniLM-L6-v2), and local Ollama (llama3.2).",
            "Core Implementation: Persistent multi-source notebooks · Hybrid Semantic + BM25 Top-K retrieval · Local Ollama chat with citations · 5-Tool Studio Forge."
        ]),
        ("05", "Introduction", [
            "BACKGROUND: Manual navigation across research files is slow; keyword-only search misses semantic context; ungrounded LLMs risk hallucination.",
            "NEED: A single persistent workspace for document ingestion, hybrid retrieval, grounded Q&A, citations, notes, and study synthesis.",
            "LOCAL-FIRST: Uses Ollama (llama3.2) at localhost:11434, local ChromaDB embeddings (all-MiniLM-L6-v2), and SQLite persistence."
        ]),
        ("06", "Problem Statement", [
            "How can a user interrogate a collection of personal research documents while keeping retrieval traceable and inference local?",
            "SCOPE: Persistent SQLite notebooks, hybrid ChromaDB + BM25 retrieval, grounded Ollama Q&A with clickable citations, and 5 Studio Forge tools.",
            "VERIFICATION: Validated with 99 pytest backend tests, 6 batch verification checks, and 10/10 live API audits."
        ]),
        ("07", "Literature Survey", [
            "1. Lewis et al. (2020) — RAG for Knowledge-Intensive NLP: Foundational RAG architecture applied locally in ORION.",
            "2. Karpukhin et al. (2020) — Dense Passage Retrieval: ORION embeds chunks with all-MiniLM-L6-v2 in ChromaDB alongside BM25.",
            "3. Guu et al. (2020) — REALM: Pre-training retriever vs. ORION's inference-time research workspace.",
            "4. Gao et al. (2023) — RAG Survey: Guides ORION's hybrid retrieval, augmentation, and grounded generation stages."
        ]),
        ("08", "Existing System", [
            "01 Traditional document search: Manual keyword matching and passage location.",
            "02 Dense retrieval pipelines: Learned vector representations before answering.",
            "03 Retrieval-augmented generation: Combining retrieved context with LLM generation.",
            "04 Cloud notebook assistants: Hosted AI workspaces requiring external cloud dependence."
        ]),
        ("09", "Issues in the Existing System", [
            "1. Manual navigation: High context-switching effort across large document sets.",
            "2. Evidence traceability: Generated answers often hide supporting source passages.",
            "3. Cloud dependency: Unsuitable for privacy-sensitive or offline local documents.",
            "4. Fragmented workflow: Search, notes, flashcards, mind maps, and audio live in separate tools."
        ]),
        ("10", "Proposed System", [
            "01 Ingest: Validate and persist sources per isolated notebook in SQLite.",
            "02 Chunk: Split content with source, section, and document-position metadata.",
            "03 Retrieve: Rank candidate chunks via Hybrid Semantic (all-MiniLM-L6-v2) + BM25 scoring.",
            "04 Generate: Grounded local Ollama (llama3.2) chat with deterministic fallback.",
            "05 Cite: Interactive citation chips linked to source viewer and passage snippets.",
            "06 Synthesize: 5-Tool Studio Forge (Audio Overview, Flashcards, Mind Map, Reports, Infographics)."
        ]),
        ("11", "Architecture Diagram", [
            "Pipeline: SOURCE LAYER (PDF/MD/TXT) -> CHUNKING (Section/Pos Meta) -> STORAGE (SQLite + ChromaDB) -> HYBRID RETRIEVAL (all-MiniLM-L6-v2 + BM25) -> GENERATION (Ollama llama3.2) -> OUTPUT (Cited Chat + 5 Studio Tools).",
            "Frontend: Next.js 14.2.23 + React 18 + TypeScript + Tailwind CSS (/api/py/:path* proxy).",
            "Backend: FastAPI + Python + SQLAlchemy 2.0 + SQLite + ChromaDB + Sentence-Transformers."
        ]),
        ("12", "Requirements", [
            "HARDWARE: Development workstation capable of running Node.js, Python, ChromaDB, and local Ollama inference.",
            "SOFTWARE: Next.js 14.2.23, React 18, TypeScript, FastAPI, SQLAlchemy 2.0, SQLite, ChromaDB, all-MiniLM-L6-v2, Ollama (llama3.2).",
            "RUNTIME: Frontend localhost:3000 · Backend API localhost:8000 · Ollama localhost:11434."
        ]),
        ("13", "Working Flow", [
            "1. Select/Create Notebook -> 2. Upload & Validate Source -> 3. Chunk + Embed in ChromaDB -> 4. Filter Active Sources",
            "5. Hybrid Top-K Retrieval (Semantic + BM25) -> 6. Build Grounded Context -> 7. Generate via Ollama llama3.2 -> 8. Navigate Citations & Studio Tools.",
            "Includes deterministic offline fallback for initial Studio generation and full SQLite artifact persistence."
        ]),
        ("14", "Results", [
            "PERSISTENCE & RAG: SQLite + ChromaDB isolation with Hybrid Semantic (all-MiniLM-L6-v2) + BM25 retrieval.",
            "CHAT & CITATIONS: Source-grounded Ollama (llama3.2) responses with safe Markdown and citation navigation.",
            "STUDIO FORGE: Persisted Audio Overview (TTS), Flashcards, Mind Map, Reports, and Infographics (SVG export).",
            "VERIFIED TESTS: 99/99 pytest backend tests, 6/6 batch4 tests, 10/10 live API audits, and clean production build."
        ]),
        ("15", "Future Enhancement", [
            "01 Cross-Encoder Reranking & OCR: Add cross-encoder reranking and complex table/OCR extraction on top of ChromaDB + BM25.",
            "02 Automated E2E & RAG Benchmarking: Add Playwright browser visual checks and precision/recall/faithfulness benchmarks.",
            "03 Local Model Training: Complete and evaluate the custom QLoRA fine-tuning pipeline for domain-specific synthesis."
        ]),
        ("16", "Conclusion", [
            "ORION unifies persistent notebook management, hybrid semantic retrieval, grounded Q&A, and 5 Studio synthesis tools locally.",
            "Built with Next.js 14, FastAPI, SQLite (SQLAlchemy 2.0), ChromaDB (all-MiniLM-L6-v2), and Ollama (llama3.2).",
            "Verified across 99 backend unit/integration tests and 10 live API audits."
        ]),
        ("17", "References", [
            "[1] P. Lewis et al., 'Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks,' arXiv:2005.11401, 2020.",
            "[2] V. Karpukhin et al., 'Dense Passage Retrieval for Open-Domain Question Answering,' EMNLP, 2020.",
            "[3] K. Guu et al., 'REALM: Retrieval-Augmented Language Model Pre-Training,' arXiv:2002.08909, 2020.",
            "[4] Y. Gao et al., 'Retrieval-Augmented Generation for Large Language Models: A Survey,' arXiv:2312.10997, 2023."
        ]),
        ("18", "Thank You", [
            "Questions are welcome.",
            "ORION · Retrieval-Augmented Document Intelligence"
        ])
    ]

    for num, title, bullets in slides_data:
        s = prs.slides.add_slide(blank)
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG
        bg.line.fill.background()

        hb = s.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(8), Inches(0.4))
        hp = hb.text_frame.paragraphs[0]
        hp.text = f"{num}  /  PROJECT ORION"
        hp.font.size = Pt(11)
        hp.font.bold = True
        hp.font.color.rgb = CYAN

        tb = s.shapes.add_textbox(Inches(0.7), Inches(0.85), Inches(11.9), Inches(0.7))
        tp = tb.text_frame.paragraphs[0]
        tp.text = title
        tp.font.size = Pt(28)
        tp.font.bold = True
        tp.font.color.rgb = WHITE

        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.75), Inches(11.9), Inches(4.9))
        card.adjustments[0] = 0.02
        card.fill.solid()
        card.fill.fore_color.rgb = CARD
        card.line.color.rgb = BORDER

        cb = s.shapes.add_textbox(Inches(1.0), Inches(2.05), Inches(11.3), Inches(4.3))
        ctf = cb.text_frame
        ctf.word_wrap = True
        for i, b in enumerate(bullets):
            p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
            p.text = f"—  {b}"
            p.font.size = Pt(16)
            p.font.color.rgb = WHITE if i == 0 else MUTED
            p.space_after = Pt(16)

    prs.save(str(OUTPUT_PATH))
    print(f"Built standalone presentation -> {OUTPUT_PATH}")

if __name__ == "__main__":
    tpl = find_template()
    if tpl:
        update_existing_template(tpl)
    else:
        build_from_scratch()
